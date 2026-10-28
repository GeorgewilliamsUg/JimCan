#!/usr/bin/env python3
"""
Generate a realistic 200-row x 10-column Excel spreadsheet of construction expenses
for Jimcan Operations Dashboard testing.
"""

from datetime import date, timedelta
from pathlib import Path
import random
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECTS = [
    'P-2026-001',  # Kira Residential Villa
    'P-2026-002',  # Najjera Commercial Arcade
    'P-2026-003',  # Entebbe Warehousing Unit
    'P-2026-004',  # Muyenga Luxury Apartments
    'OVERHEAD',     # Company Head Office & Fleet
]

CATEGORIES = [
    'Direct Cost - Materials',
    'Direct Cost - Labour',
    'Direct Cost - Equipment Hire',
    'Direct Cost - Subcontract',
    'Site Overhead - Fuel & Transport',
    'Site Overhead - Security & Safety',
    'Administrative Overhead',
]

PAYMENT_METHODS = ['Mobile Money', 'Bank', 'Cash', 'Cheque']

SUPPLIERS = {
    'Direct Cost - Materials': [
        'Tororo Cement Depot', 'Hima Cement Uganda', 'Kalyango Hardware Ltd',
        'Roofings Middle East Steel', 'Uganda Clays Kajjansi', 'Mukwano Timber Traders',
        'Crown Paints Uganda', 'Chint Electrical Supplies', 'Hardware City Nakasero'
    ],
    'Direct Cost - Labour': [
        'Kira Site Casual Team (12 Pax)', 'Masonry Gang Lead By Sula',
        'Steel Fixers Sub-Gang', 'Formwork Carpenters Union', 'Plumbing Tech Team'
    ],
    'Direct Cost - Equipment Hire': [
        'Kato Plant & Earthmovers', 'Mukono Compactor Services', 'Kampala Mixer Rentals',
        'Scaffolding Masters Uganda'
    ],
    'Direct Cost - Subcontract': [
        'Apex Structural Glazing', 'Solid Rock Geo-Drillers', 'Pro-Tech Electricals',
        'Waterproof Solutions Ltd'
    ],
    'Site Overhead - Fuel & Transport': [
        'TotalEnergies Mukono Bypass', 'Shell Kira Road Station', 'Stabex Naalya Express',
        'Isuzu Tipper Haulage'
    ],
    'Site Overhead - Security & Safety': [
        'Delta Guard Security Services', 'Protector Safety Gear Ltd', 'Site First Aid & Water Supply'
    ],
    'Administrative Overhead': [
        'Airtel Uganda Corporate Internet', 'Umeme Power Utility', 'National Water & Sewerage Corp',
        'Office Stationery Hub'
    ]
}

SAMPLE_DESCRIPTIONS = {
    'Direct Cost - Materials': [
        '50 bags 42.5N Pozzolana Cement', '10-wheeler tipper lake sand delivery',
        '15-tonne crushed aggregate (20mm)', '20 lengths high-tensile rebar 12mm',
        '30 lengths deformed steel bar 10mm', '100 pieces solid precast concrete blocks (200mm)',
        '40 sheets BRC mesh A142', '5 bundles hardwood timber (4x2)',
        '20 pieces marine plywood formwork (18mm)', '4 drums water-repellent concrete admixture',
        '5 buckets vinyl matt emulsion paint (20L)', '100m PVC conduit pipes (25mm)'
    ],
    'Direct Cost - Labour': [
        'Weekly casual labor wages for foundation excavation',
        'Masonry team day rate for ground floor walling',
        'Steel fixing sub-gang milestone disbursement',
        'Carpenters day rate for suspended slab formwork',
        'Overtime night pour concrete placing allowance'
    ],
    'Direct Cost - Equipment Hire': [
        '8-hour excavator hire with operator for basement digging',
        '3-day diesel concrete mixer hire on site',
        '2-day plate compactor hire for hardcore preparation',
        'Mobile crane half-day hire for steel beam lifting'
    ],
    'Direct Cost - Subcontract': [
        'Phase 1 structural borehole core drilling advance',
        'Aluminum sliding window fabrication milestone 1',
        'Sub-circuit wiring and distribution board installation',
        'Foundation basement bituminous waterproofing application'
    ],
    'Site Overhead - Fuel & Transport': [
        '100 Litres diesel for site standby generator',
        '50 Litres diesel for JCB backhoe operation',
        'Delivery transport charges for timber offloading',
        'Material runner motorcycle fuel reimbursement'
    ],
    'Site Overhead - Security & Safety': [
        'Monthly perimeter night guard service on site',
        '20 pairs heavy-duty steel toe safety boots',
        'Reflective high-visibility site vests and safety helmets',
        'Weekly bottled drinking water dispenser refills for workers'
    ],
    'Administrative Overhead': [
        'Head office fiber internet subscription for October',
        'Site office prepaid Umeme electricity tokens',
        'Site office printer paper and printing supplies',
        'Monthly legal and site permit compliance fees'
    ]
}


def generate_expenses_workbook(output_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Site Expenses"

    # Define Palette & Styles
    brand_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Charcoal slate
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    money_font = Font(name="Calibri", size=10, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    headers = [
        "Date",
        "Project Code",
        "Category",
        "Voucher / Ref",
        "Description",
        "Net (UGX)",
        "VAT (UGX)",
        "Total Gross (UGX)",
        "Paid To",
        "Payment Method"
    ]

    ws.append(headers)
    ws.row_dimensions[1].height = 26

    # Style Header Row
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = brand_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_thin

    start_date = date(2026, 5, 1)
    end_date = date(2026, 10, 26)
    delta_days = (end_date - start_date).days

    random.seed(42)  # Deterministic test set

    for row_num in range(2, 202):
        # 1. Date
        cur_date = start_date + timedelta(days=random.randint(0, delta_days))
        date_str = cur_date.strftime("%Y-%m-%d")

        # 2. Category & Project
        cat = random.choice(CATEGORIES)
        proj = 'OVERHEAD' if 'Administrative' in cat else random.choice(PROJECTS)

        # 3. Voucher
        voucher = f"PV-2026-{row_num - 1:04d}"

        # 4. Description & Paid To
        desc = random.choice(SAMPLE_DESCRIPTIONS[cat])
        supplier = random.choice(SUPPLIERS[cat])

        # 5. Financial figures
        if 'Materials' in cat:
            net = round(random.uniform(350000, 4800000), -3)
            vat_applies = random.random() > 0.35  # ~65% VAT formal
        elif 'Subcontract' in cat:
            net = round(random.uniform(1200000, 6500000), -3)
            vat_applies = True
        elif 'Labour' in cat:
            net = round(random.uniform(180000, 1500000), -3)
            vat_applies = False  # Direct informal/wage labor exempt from VAT
        elif 'Equipment' in cat:
            net = round(random.uniform(400000, 2200000), -3)
            vat_applies = random.random() > 0.5
        elif 'Fuel' in cat:
            net = round(random.uniform(150000, 850000), -3)
            vat_applies = True
        elif 'Security' in cat:
            net = round(random.uniform(250000, 900000), -3)
            vat_applies = True
        else:  # Administrative
            net = round(random.uniform(80000, 1200000), -3)
            vat_applies = random.random() > 0.4

        vat = round(net * 0.18, -2) if vat_applies else 0.0
        gross = net + vat

        # 6. Payment Method
        if 'Labour' in cat or net < 300000:
            method = random.choice(['Mobile Money', 'Cash'])
        elif gross > 3000000:
            method = random.choice(['Bank', 'Cheque'])
        else:
            method = random.choice(PAYMENT_METHODS)

        row_data = [
            date_str,
            proj,
            cat,
            voucher,
            desc,
            net,
            vat,
            gross,
            supplier,
            method
        ]
        ws.append(row_data)
        ws.row_dimensions[row_num].height = 20

        # Style data cells
        for col_idx in range(1, 11):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = border_thin

            if col_idx in (1, 2, 4, 10):
                cell.alignment = align_center
            elif col_idx in (6, 7, 8):
                cell.alignment = align_right
                cell.number_format = '#,##0'
                if col_idx == 8:
                    cell.font = money_font
            else:
                cell.alignment = align_left

    # Total Row at the bottom
    tot_row = 202
    ws.row_dimensions[tot_row].height = 24
    ws.cell(row=tot_row, column=5, value="TOTALS (UGX)").font = Font(name="Calibri", size=11, bold=True)
    ws.cell(row=tot_row, column=5).alignment = align_right

    for c_idx in (6, 7, 8):
        c_letter = get_column_letter(c_idx)
        t_cell = ws.cell(row=tot_row, column=c_idx)
        t_cell.value = f"=SUM({c_letter}2:{c_letter}201)"
        t_cell.font = Font(name="Calibri", size=11, bold=True)
        t_cell.number_format = '#,##0'
        t_cell.alignment = align_right
        t_cell.border = Border(top=Side(style='thin'), bottom=Side(style='double'))

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.number_format == '#,##0' and isinstance(cell.value, (int, float)):
                val_str = f"{int(cell.value):,}"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_path)
    print(f"Successfully generated 200 rows of test expenses at: {output_path}")


if __name__ == "__main__":
    out_file = Path(__file__).resolve().parent.parent / "expenses_test_data_200.xlsx"
    generate_expenses_workbook(out_file)
